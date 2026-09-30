"""
=============================================================================
  CHAMTN EXCEL REPORTER (4 SHEETS)
  Tạo báo cáo kết quả thi xuất bản chuẩn Bộ GD&ĐT 4 sheets:
    Sheet 1: "Tổng hợp"       - Bảng điểm danh sách học sinh, điểm thành phần & tổng điểm.
    Sheet 2: "Bảng điểm"      - Điểm chi tiết từng câu trắc nghiệm của từng học sinh.
    Sheet 3: "Bài làm chi tiết" - Đáp án học sinh đã chọn (A/B/C/D, Đ/S, số điền).
    Sheet 4: "Phân tích câu hỏi" - Thống kê độ khó, tỉ lệ đúng, phương án nhiễu & độ phân cách.
=============================================================================
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import numpy as np
from typing import List, Dict, Any

# Bảng màu chuyên nghiệp
HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Xanh navy đậm
SUB_HEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid") # Xanh dương
LIGHT_BLUE_FILL = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
CORRECT_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")   # Xanh lá nhạt
WRONG_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")     # Đỏ nhạt
BLANK_FILL = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")     # Vàng nhạt

WHITE_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TEXT_BOLD = Font(name="Calibri", size=11, bold=True, color="000000")
TEXT_REGULAR = Font(name="Calibri", size=11, color="000000")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1E3A8A")

THIN_SIDE = Side(border_style="thin", color="CBD5E1")
BORDER_ALL = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def export_exam_results_excel(
    exam_info: Dict[str, Any],
    results: List[Dict[str, Any]],
    output_path: str
) -> str:
    """
    Xuất file Excel báo cáo kỳ thi 4 sheets chuẩn ChamTN.
    """
    wb = openpyxl.Workbook()
    # Xóa sheet mặc định
    wb.remove(wb.active)

    exam_title = exam_info.get("name", "BÁO CÁO KẾT QUẢ KỲ THI")
    exam_subject = exam_info.get("subject", "Trắc nghiệm")
    exam_date = exam_info.get("date", "")
    n_p1 = exam_info.get("counts", {}).get("p1", 40)
    n_p2 = exam_info.get("counts", {}).get("p2", 8)
    n_p3 = exam_info.get("counts", {}).get("p3", 6)

    # =========================================================================
    # SHEET 1: TỔNG HỢP
    # =========================================================================
    ws1 = wb.create_sheet(title="Tổng hợp")
    ws1.views.sheetView[0].showGridLines = True

    # Tiêu đề báo cáo
    ws1.merge_cells("A1:J1")
    cell_t = ws1["A1"]
    cell_t.value = exam_title.upper()
    cell_t.font = TITLE_FONT
    cell_t.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 35

    ws1.merge_cells("A2:J2")
    ws1["A2"].value = f"Môn thi: {exam_subject}  |  Ngày thi: {exam_date}  |  Tổng số bài: {len(results)}"
    ws1["A2"].font = Font(name="Calibri", size=11, italic=True, color="475569")
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    # Tiêu đề bảng
    headers1 = [
        "STT", "Số báo danh", "Họ và tên", "Lớp", "Phòng thi",
        "Mã đề", "Điểm Phần I", "Điểm Phần II", "Điểm Phần III", "Tổng điểm"
    ]
    ws1.append([])  # Row 3 trống
    ws1.append(headers1)  # Row 4 headers
    ws1.row_dimensions[4].height = 28

    for col_idx in range(1, len(headers1) + 1):
        c = ws1.cell(row=4, column=col_idx)
        c.fill = HEADER_FILL
        c.font = WHITE_BOLD
        c.alignment = ALIGN_CENTER
        c.border = BORDER_ALL

    # Dữ liệu học sinh
    for idx, r in enumerate(results, 1):
        row_num = 4 + idx
        ws1.row_dimensions[row_num].height = 22
        ws1.append([
            idx,
            r.get("sbd", ""),
            r.get("name", "Thí sinh " + str(idx)),
            r.get("class_name", ""),
            r.get("room", ""),
            r.get("made", ""),
            r.get("p1_score", 0.0),
            r.get("p2_score", 0.0),
            r.get("p3_score", 0.0),
            r.get("total_score", 0.0)
        ])
        for col_idx in range(1, len(headers1) + 1):
            c = ws1.cell(row=row_num, column=col_idx)
            c.font = TEXT_BOLD if col_idx == 10 else TEXT_REGULAR
            c.border = BORDER_ALL
            if col_idx in [1, 2, 4, 5, 6]:
                c.alignment = ALIGN_CENTER
            elif col_idx in [7, 8, 9, 10]:
                c.alignment = ALIGN_RIGHT
                c.number_format = "0.00"
                if col_idx == 10:
                    c.fill = LIGHT_BLUE_FILL

    # =========================================================================
    # SHEET 2: BẢNG ĐIỂM CHI TIẾT
    # =========================================================================
    ws2 = wb.create_sheet(title="Bảng điểm")
    ws2.views.sheetView[0].showGridLines = True

    headers2 = ["STT", "SBD", "Họ và tên", "Mã đề"]
    for q in range(1, n_p1 + 1):
        headers2.append(f"C{q}")
    for q in range(1, n_p2 + 1):
        headers2.append(f"II.{q}")
    for q in range(1, n_p3 + 1):
        headers2.append(f"III.{q}")
    headers2.append("Tổng điểm")

    ws2.append(headers2)
    ws2.row_dimensions[1].height = 26
    for col_idx in range(1, len(headers2) + 1):
        c = ws2.cell(row=1, column=col_idx)
        c.fill = SUB_HEADER_FILL
        c.font = WHITE_BOLD
        c.alignment = ALIGN_CENTER
        c.border = BORDER_ALL

    for idx, r in enumerate(results, 1):
        row_data = [idx, r.get("sbd", ""), r.get("name", ""), r.get("made", "")]
        p1_det = r.get("p1_details", {})
        for q in range(1, n_p1 + 1):
            row_data.append(p1_det.get(q, {}).get("pts", 0.0))
        p2_det = r.get("p2_details", {})
        for q in range(1, n_p2 + 1):
            row_data.append(p2_det.get(q, {}).get("pts", 0.0))
        p3_det = r.get("p3_details", {})
        for q in range(1, n_p3 + 1):
            row_data.append(p3_det.get(q, {}).get("pts", 0.0))
        row_data.append(r.get("total_score", 0.0))

        ws2.append(row_data)
        row_num = 1 + idx
        for col_idx in range(1, len(headers2) + 1):
            c = ws2.cell(row=row_num, column=col_idx)
            c.border = BORDER_ALL
            c.alignment = ALIGN_CENTER
            if col_idx > 4:
                c.number_format = "0.00"

    # =========================================================================
    # SHEET 3: BÀI LÀM CHI TIẾT (ĐÁP ÁN HỌC SINH ĐÃ CHỌN)
    # =========================================================================
    ws3 = wb.create_sheet(title="Bài làm chi tiết")
    ws3.views.sheetView[0].showGridLines = True

    headers3 = ["STT", "SBD", "Họ và tên", "Mã đề"]
    for q in range(1, n_p1 + 1):
        headers3.append(f"C{q}")
    for q in range(1, n_p2 + 1):
        headers3.append(f"II.{q}")
    for q in range(1, n_p3 + 1):
        headers3.append(f"III.{q}")

    ws3.append(headers3)
    ws3.row_dimensions[1].height = 26
    for col_idx in range(1, len(headers3) + 1):
        c = ws3.cell(row=1, column=col_idx)
        c.fill = HEADER_FILL
        c.font = WHITE_BOLD
        c.alignment = ALIGN_CENTER
        c.border = BORDER_ALL

    for idx, r in enumerate(results, 1):
        row_data = [idx, r.get("sbd", ""), r.get("name", ""), r.get("made", "")]
        p1_det = r.get("p1_details", {})
        for q in range(1, n_p1 + 1):
            row_data.append(p1_det.get(q, {}).get("got", "-"))

        p2_det = r.get("p2_details", {})
        for q in range(1, n_p2 + 1):
            subs = p2_det.get(q, {}).get("subs", {})
            p2_str = "".join([subs.get(s, {}).get("got", "-") for s in ["a", "b", "c", "d"]])
            row_data.append(p2_str if p2_str else "-")

        p3_det = r.get("p3_details", {})
        for q in range(1, n_p3 + 1):
            row_data.append(p3_det.get(q, {}).get("got", "-"))

        ws3.append(row_data)
        row_num = 1 + idx
        for col_idx in range(1, len(headers3) + 1):
            c = ws3.cell(row=row_num, column=col_idx)
            c.border = BORDER_ALL
            c.alignment = ALIGN_CENTER
            # Bôi màu: đúng màu xanh, sai màu đỏ, trống màu vàng
            if col_idx > 4:
                val = str(c.value)
                if col_idx <= 4 + n_p1:
                    q_num = col_idx - 4
                    is_corr = p1_det.get(q_num, {}).get("is_correct", False)
                    if not val or val == "-":
                        c.fill = BLANK_FILL
                    elif is_corr:
                        c.fill = CORRECT_FILL
                    else:
                        c.fill = WRONG_FILL

    # =========================================================================
    # SHEET 4: PHÂN TÍCH CÂU HỎI (STATISTICAL ITEM ANALYSIS)
    # =========================================================================
    ws4 = wb.create_sheet(title="Phân tích câu hỏi")
    ws4.views.sheetView[0].showGridLines = True

    headers4 = [
        "Câu hỏi", "Phần", "Đáp án chuẩn", "Số bài làm",
        "Số câu đúng", "Tỉ lệ đúng (%)", "Tỉ lệ chọn A (%)",
        "Tỉ lệ chọn B (%)", "Tỉ lệ chọn C (%)", "Tỉ lệ chọn D (%)",
        "Tỉ lệ bỏ trống (%)", "Độ phân cách (D)"
    ]
    ws4.append(headers4)
    ws4.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers4) + 1):
        c = ws4.cell(row=1, column=col_idx)
        c.fill = HEADER_FILL
        c.font = WHITE_BOLD
        c.alignment = ALIGN_CENTER
        c.border = BORDER_ALL

    n_total = max(1, len(results))

    # Xếp hạng học sinh để tính Độ phân cách (Top 27% vs Bottom 27%)
    sorted_results = sorted(results, key=lambda x: x.get("total_score", 0), reverse=True)
    top_group_size = max(1, int(len(sorted_results) * 0.27))
    top_group = sorted_results[:top_group_size]
    bot_group = sorted_results[-top_group_size:]

    for q in range(1, n_p1 + 1):
        want_ans = ""
        correct_count = 0
        blank_count = 0
        opt_counts = {"A": 0, "B": 0, "C": 0, "D": 0}

        for r in results:
            det = r.get("p1_details", {}).get(q, {})
            want_ans = det.get("want", want_ans)
            got = det.get("got", "")
            if det.get("is_correct", False):
                correct_count += 1
            if not got or got == "-":
                blank_count += 1
            elif got in opt_counts:
                opt_counts[got] += 1

        top_corr = sum(1 for r in top_group if r.get("p1_details", {}).get(q, {}).get("is_correct", False))
        bot_corr = sum(1 for r in bot_group if r.get("p1_details", {}).get(q, {}).get("is_correct", False))
        disc_index = round((top_corr - bot_corr) / float(top_group_size), 2)

        p_correct = round(correct_count * 100.0 / n_total, 1)
        p_blank = round(blank_count * 100.0 / n_total, 1)

        ws4.append([
            f"Câu {q}",
            "Phần I",
            want_ans,
            n_total,
            correct_count,
            f"{p_correct}%",
            f"{round(opt_counts['A'] * 100.0 / n_total, 1)}%",
            f"{round(opt_counts['B'] * 100.0 / n_total, 1)}%",
            f"{round(opt_counts['C'] * 100.0 / n_total, 1)}%",
            f"{round(opt_counts['D'] * 100.0 / n_total, 1)}%",
            f"{p_blank}%",
            disc_index
        ])

        row_idx = 1 + q
        for col_idx in range(1, len(headers4) + 1):
            c = ws4.cell(row=row_idx, column=col_idx)
            c.border = BORDER_ALL
            c.alignment = ALIGN_CENTER

    # Tự động điều chỉnh độ rộng các cột cho tất cả các sheet
    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    val_str = str(cell.value)
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    wb.save(output_path)
    return output_path
