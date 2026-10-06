"""Read bounded spreadsheet values only; preview is confirmed with a signed payload."""
import csv
import io
from itertools import islice
from zipfile import ZipFile, BadZipFile

from django.core import signing
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

MAX_ROWS = 2000
IMPORT_SALT = 'gradeflow.classroom-import.v1'


def read_table(upload):
    if upload.size > 5 * 1024 * 1024:
        raise ValidationError('File tối đa 5 MB.')
    raw = upload.read()
    if upload.name.lower().endswith('.csv'):
        try:
            decoded = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ValidationError('Lưu CSV theo mã UTF-8 để giữ tiếng Việt.')
        try:
            dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        values = list(islice(csv.reader(io.StringIO(decoded), dialect), MAX_ROWS + 2))
    elif upload.name.lower().endswith('.xlsx'):
        try:
            with ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValidationError('File Excel quá lớn sau khi giải nén.')
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            values = []
            for row in workbook.active.iter_rows():
                cells = []
                for cell in row[:50]:
                    value = cell.value
                    if cell.data_type == 'f':
                        raise ValidationError('Danh sách cần giá trị trực tiếp, không dùng công thức.')
                    # Excel numeric cells with 000000 formatting retain their displayed SBD.
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        if float(value).is_integer():
                            value = str(int(value))
                            fmt = cell.number_format
                            if fmt and set(fmt) == {'0'}:
                                value = value.zfill(len(fmt))
                    cells.append('' if value is None else str(value).strip())
                values.append(cells)
                if len(values) > MAX_ROWS + 1:
                    break
            workbook.close()
        except ValidationError:
            raise
        except (BadZipFile, ValueError, OSError, KeyError, TypeError):
            raise ValidationError('Không đọc được file Excel. Chọn file .xlsx hợp lệ.')
    else:
        raise ValidationError('Chọn file Excel .xlsx hoặc CSV.')
    values = [row for row in values if any(str(c).strip() for c in row)]
    if len(values) < 2:
        raise ValidationError('File cần hàng tiêu đề và ít nhất một học sinh.')
    if len(values) > MAX_ROWS + 1:
        raise ValidationError(f'Mỗi lần nhập tối đa {MAX_ROWS} học sinh.')
    return [str(c)[:100] for c in values[0][:50]], values[1:]


def preview_rows(classroom, headers, values, name_column, id_column):
    if name_column == id_column or min(name_column, id_column) < 0 or max(name_column, id_column) >= len(headers):
        raise ValidationError('Chọn hai cột khác nhau cho họ tên và số báo danh.')
    rows, issues, seen = [], [], set()
    existing = {s.student_id: s for s in classroom.students.all()}
    for index, cells in enumerate(values, 2):
        name = str(cells[name_column]).strip() if name_column < len(cells) else ''
        sbd = str(cells[id_column]).strip() if id_column < len(cells) else ''
        error = ''
        if not name or not sbd:
            error = 'Thiếu họ tên hoặc số báo danh.'
        elif len(name) > 200 or len(sbd) > 50:
            error = 'Họ tên hoặc số báo danh quá dài.'
        elif sbd in seen:
            error = 'Số báo danh trùng trong file.'
        elif sbd in existing and existing[sbd].name != name:
            error = 'Số báo danh đã thuộc học sinh khác trong lớp.'
        seen.add(sbd)
        action = 'existing' if sbd in existing else 'add'
        rows.append({'row': index, 'name': name, 'student_id': sbd, 'action': action, 'error': error})
        if error:
            issues.append({'row': index, 'message': error})
    token = signing.dumps({'owner': classroom.owner_id, 'classroom': classroom.pk,
                           'rows': [{'name': r['name'], 'student_id': r['student_id']} for r in rows]},
                          salt=IMPORT_SALT, compress=True) if not issues else None
    return {'rows': rows, 'issues': issues, 'confirmation': token,
            'add_count': sum(r['action'] == 'add' for r in rows),
            'existing_count': sum(r['action'] == 'existing' for r in rows)}
