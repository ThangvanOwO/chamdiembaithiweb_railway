"""Read-only reporting of saved grades. Never invokes the OMR/scoring engine."""
import csv
import io
import json
import math
from collections import defaultdict
from pathlib import Path
from statistics import mean, median
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from grading.models import Submission
from .models import ExamRosterEntry


def object_json(raw):
    try:
        result = json.loads(raw) if isinstance(raw, str) else raw
        return result if isinstance(result, dict) else {}
    except (TypeError, ValueError):
        return {}


def unclear(value):
    return value is None or str(value).strip().upper() in {'', 'X', 'NULL'} or '?' in str(value)


def recognition_issues(sub):
    issues = []
    answers = object_json(sub.answers_detected)
    for part in (1, 2, 3):
        for question, value in object_json(answers.get(f'part{part}')).items():
            if isinstance(value, dict):
                issues.extend(f'Phần {part} · Câu {question}{key}' for key, v in value.items() if unclear(v))
            elif unclear(value):
                issues.append(f'Phần {part} · Câu {question}')
    if sub.status != 'completed':
        issues.append('Bài lỗi' if sub.status == 'error' else 'Bài chưa chấm xong')
    if not sub.student_id.strip() or '?' in sub.student_id:
        issues.append('Số báo danh chưa rõ')
    return issues


def review_of(sub):
    try:
        return sub.toolbox_review
    except Submission.toolbox_review.RelatedObjectDoesNotExist:
        return None


def representative(subs):
    completed = [s for s in subs if s.status == 'completed' and s.score is not None and math.isfinite(s.score)]
    selected = [s for s in completed if review_of(s) and review_of(s).representative]
    # Stable even when imports/older rows share timestamps.
    return max(selected or completed or subs, key=lambda s: (s.uploaded_at, s.pk)) if subs else None


def report_data(exam, classroom_id=None):
    roster = list(ExamRosterEntry.objects.filter(exam=exam, active=True))
    by_id = {r.pk: r for r in roster}
    by_sbd = {r.student_id: r for r in roster}
    submissions = list(Submission.objects.filter(exam=exam, teacher_id=exam.teacher_id)
                       .select_related('variant', 'toolbox_review').order_by('-uploaded_at', '-id'))
    groups = defaultdict(list)
    for sub in submissions:
        review = review_of(sub)
        entry = by_id.get(review.roster_entry_id) if review and review.roster_entry_id else by_sbd.get(sub.student_id.strip())
        raw_sbd = sub.student_id.strip()
        key = ('roster', entry.pk) if entry else ('unknown', raw_sbd if raw_sbd and '?' not in raw_sbd else f'#{sub.pk}')
        groups[key].append(sub)
    rows, chosen = [], []

    def append_row(entry, subs):
        sub = representative(subs)
        review = review_of(sub) if sub else None
        issues = recognition_issues(sub) if sub else []
        if sub and roster and not entry:
            issues.append('Chưa ghép học sinh')
        if len(subs) > 1:
            issues.append(f'{len(subs)} phiếu cùng số báo danh')
        scored = sub and sub.status == 'completed' and sub.score is not None and math.isfinite(sub.score)
        rows.append({'roster_id': entry.pk if entry else None,
                     'student_id': entry.student_id if entry else sub.student_id,
                     'name': entry.student_name if entry else sub.student_name,
                     'class_name': entry.class_name if entry else '',
                     'classroom_id': entry.classroom_id if entry else None,
                     'score': sub.score_10 if scored else None,
                     'status': 'Chưa có bài' if not sub else 'Cần kiểm tra' if issues and not (review and review.reviewed_at)
                     else 'Đã có bài' if scored else 'Chờ chấm / lỗi',
                     'submission_id': sub.pk if sub else None,
                     'variant_code': sub.variant.variant_code if sub and sub.variant else '',
                     'issues': issues, 'reviewed': bool(review and review.reviewed_at),
                     'duplicate_count': len(subs), 'candidates': [s.pk for s in subs]})
        if scored:
            chosen.append(sub)

    for entry in roster:
        if classroom_id is None or entry.classroom_id == classroom_id:
            append_row(entry, groups.get(('roster', entry.pk), []))
    if classroom_id is None:
        for key, subs in groups.items():
            if key[0] == 'unknown':
                append_row(None, subs)
    scores = [row['score'] for row in rows if row['score'] is not None]
    histogram = [0] * 10
    for score in scores:
        histogram[min(9, max(0, int(score)))] += 1

    items = defaultdict(lambda: {'correct': 0, 'wrong': 0, 'unclear': 0})
    missing = 0
    for sub in chosen:
        details = object_json(sub.detail_json)
        variant = sub.variant.variant_code if sub.variant else ''
        found = False

        def observe(part, question, leaf, subquestion=''):
            nonlocal found
            if not isinstance(leaf, dict) or not isinstance(leaf.get('is_correct'), bool):
                return
            found = True
            item = items[(variant, part, str(question), str(subquestion))]
            state = 'unclear' if unclear(leaf.get('student')) else 'correct' if leaf['is_correct'] else 'wrong'
            item[state] += 1

        for part in (1, 2, 3):
            for question, leaf in object_json(details.get(f'part{part}_detail')).items():
                if part == 2 and isinstance(leaf, dict):
                    leaves = [leaf.get(k) for k in ('a', 'b', 'c', 'd')]
                    for key in ('a', 'b', 'c', 'd'):
                        observe(part, question, leaf.get(key), key)
                    if all(isinstance(v, dict) and isinstance(v.get('is_correct'), bool) for v in leaves):
                        observe(part, question, {'student': '?' if any(unclear(v.get('student')) for v in leaves) else 'known',
                                                 'is_correct': all(v['is_correct'] for v in leaves)})
                else:
                    observe(part, question, leaf)
        if not found:
            missing += 1
    analysis = []
    def question_order(item):
        variant, part, question, subquestion = item[0]
        number = (0, int(question)) if question.isdigit() else (1, question)
        return variant, part, number, subquestion
    for (variant, part, question, subquestion), counts in sorted(items.items(), key=question_order):
        known = counts['correct'] + counts['wrong']
        analysis.append({'variant_code': variant, 'part': part, 'question': question, 'subquestion': subquestion,
                         **counts, 'known_count': known,
                         'correct_rate': round(100 * counts['correct'] / known, 2) if known else None})
    classes = {r.classroom_id: r.class_name for r in roster if r.classroom_id}
    return {'exam_id': exam.pk, 'exam_title': exam.title, 'classroom_id': classroom_id, 'rows': rows,
            'classes': [{'id': pk, 'name': name} for pk, name in classes.items()],
            'summary': {'scored': len(scores), 'students': len(rows), 'missing': sum(r['submission_id'] is None for r in rows),
                        'needs_review': sum(r['status'] == 'Cần kiểm tra' for r in rows),
                        'mean': round(mean(scores), 2) if scores else None,
                        'median': round(median(scores), 2) if scores else None,
                        'min': min(scores) if scores else None, 'max': max(scores) if scores else None,
                        'histogram': histogram, 'missing_detail': missing}, 'questions': analysis}


def grade_rows(data):
    return [['Số báo danh', 'Họ tên', 'Lớp', 'Mã đề', 'Điểm', 'Trạng thái', 'Ghi chú']] + [
        [r['student_id'], r['name'], r['class_name'], r['variant_code'], r['score'], r['status'], '; '.join(r['issues'])]
        for r in data['rows']]


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def csv_report(data):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerows([[safe_cell(v) if v is not None else '' for v in row] for row in grade_rows(data)])
    return out.getvalue().encode('utf-8-sig')


def xlsx_report(data):
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = data['summary']
    tables = {
        'Bảng điểm': grade_rows(data),
        'Trạng thái học sinh': [['Số báo danh', 'Họ tên', 'Lớp', 'Trạng thái', 'Đã đối chiếu']] + [
            [r['student_id'], r['name'], r['class_name'], r['status'], 'Có' if r['reviewed'] else 'Chưa'] for r in data['rows']],
        'Tổng hợp': [['Nội dung', 'Giá trị'], ['Bài có điểm', summary['scored']], ['Chưa có bài', summary['missing']],
                    ['Cần kiểm tra', summary['needs_review']], ['Trung bình', summary['mean']], ['Trung vị', summary['median']],
                    ['Thấp nhất', summary['min']], ['Cao nhất', summary['max']], ['Thiếu chi tiết', summary['missing_detail']]],
        'Thống kê câu': [['Mã đề', 'Phần', 'Câu', 'Ý', 'Đúng', 'Sai', 'Chưa rõ', 'Tỷ lệ đúng trên ô rõ (%)']] + [
            [q['variant_code'], q['part'], q['question'], q['subquestion'], q['correct'], q['wrong'], q['unclear'], q['correct_rate']]
            for q in data['questions']],
    }
    for title, table in tables.items():
        sheet = workbook.create_sheet(title)
        for row in table:
            sheet.append([safe_cell(v) for v in row])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='087E79')
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cells in sheet.columns:
            sheet.column_dimensions[cells[0].column_letter].width = min(55, max(14, max(len(str(c.value or '')) for c in cells) + 2))
        # SBD is explicitly text, including numeric-looking leading zeros.
        if title in ('Bảng điểm', 'Trạng thái học sinh'):
            for row in sheet.iter_rows(min_row=2):
                row[0].number_format = '@'
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def pdf_font():
    if 'GradeFlow' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('GradeFlow', str(Path(__file__).parent / 'assets/Manrope-Regular.ttf')))
    return 'GradeFlow'


def pdf_report(data):
    stream = io.BytesIO()
    font = pdf_font()
    style = ParagraphStyle('gf', fontName=font, fontSize=8, leading=12)
    title_style = ParagraphStyle('title', parent=style, fontSize=16, leading=23)
    paragraph = lambda text: Paragraph(escape(str(text if text is not None else '')), style)
    summary = data['summary']
    story = [Paragraph(escape(f"GradeFlow · {data['exam_title']}"), title_style), Spacer(1, 14),
             paragraph(f"Bài có điểm: {summary['scored']} · Chưa có bài: {summary['missing']} · Cần kiểm tra: {summary['needs_review']}"),
             paragraph(f"Trung bình: {summary['mean']} · Trung vị: {summary['median']} · Thấp nhất: {summary['min']} · Cao nhất: {summary['max']}"),
             Spacer(1, 12)]
    table = Table([[paragraph(c) for c in row] for row in grade_rows(data)],
                  colWidths=[65, 120, 65, 45, 40, 90, 330 - 20], repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#dcefea')),
                              ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                              ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#cadbd7'))]))
    story.append(table)
    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.drawString(24, 16, f'GradeFlow · Trang {doc.page}')
    SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=24, rightMargin=24,
                      topMargin=24, bottomMargin=30).build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()
