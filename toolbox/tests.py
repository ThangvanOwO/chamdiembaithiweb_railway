import io
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from grading.models import Exam, ExamVariant, Submission
from .models import Classroom, Student, ExamRosterEntry, SupportTicket
from .reports import report_data, xlsx_report, csv_report, pdf_report


class ToolboxTests(TestCase):
    def setUp(self):
        self.owner = get_user_model().objects.create_user('toolbox_teacher', password='test-password')
        self.other = get_user_model().objects.create_user('other_teacher', password='test-password')
        self.client = APIClient()
        self.client.force_authenticate(self.owner)
        self.exam = Exam.objects.create(teacher=self.owner, title='Toán', num_questions=2)
        self.room = Classroom.objects.create(owner=self.owner, name='10A', school_year='2026–2027')

    def student(self, sbd='000123', name='Nguyễn An'):
        return Student.objects.create(classroom=self.room, student_id=sbd, name=name)

    def submission(self, sbd='000123', score=7, **kwargs):
        return Submission.objects.create(teacher=self.owner, exam=self.exam, student_id=sbd, score=score,
                                         status='completed', **kwargs)

    def attach(self):
        return self.client.post(f'/api/v1/exams/{self.exam.pk}/roster/', {'classroom_ids': [self.room.pk]}, format='json')

    def test_ownership_all_routes_and_anonymous(self):
        foreign = Classroom.objects.create(owner=self.other, name='Private')
        self.assertEqual(self.client.get(f'/api/v1/classrooms/{foreign.pk}/').status_code, 404)
        for path in ('report/', 'export/xlsx/', 'roster/'):
            foreign_exam = Exam.objects.create(teacher=self.other, title='Private')
            self.assertEqual(self.client.get(f'/api/v1/exams/{foreign_exam.pk}/{path}').status_code, 404)
        self.client.force_authenticate(None)
        self.assertIn(self.client.get('/api/v1/classrooms/').status_code, (401, 403))

    def test_snapshot_and_duplicate_sbd_rollback(self):
        student = self.student()
        self.assertEqual(self.attach().status_code, 200)
        student.name = 'Tên đã đổi'
        student.save()
        self.room.name = 'Lớp đã đổi'
        self.room.save()
        entry = ExamRosterEntry.objects.get(exam=self.exam)
        self.assertEqual(entry.student_name, 'Nguyễn An')
        self.assertEqual(entry.class_name, '10A')
        room2 = Classroom.objects.create(owner=self.owner, name='10B')
        Student.objects.create(classroom=room2, student_id='000123', name='Người khác')
        response = self.client.post(f'/api/v1/exams/{self.exam.pk}/roster/', {'classroom_ids': [room2.pk]}, format='json')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(ExamRosterEntry.objects.count(), 1)

    def test_import_excel_format_mapping_confirmation_replay(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['Họ tên', 'SBD'])
        sheet.append(['Nguyễn An', 123])
        sheet['B2'].number_format = '000000'
        raw = io.BytesIO()
        book.save(raw)
        url = f'/api/v1/classrooms/{self.room.pk}/import/'
        first = self.client.post(url + 'preview/', {'file': SimpleUploadedFile('lop.xlsx', raw.getvalue())}, format='multipart')
        self.assertEqual(first.status_code, 200)
        self.assertEqual(first.data['sample'][0][1], '000123')
        preview = self.client.post(url + 'preview/', {'source_token': first.data['source_token'],
                                                     'name_column': 0, 'id_column': 1}, format='json')
        self.assertEqual(preview.status_code, 200)
        confirmed = self.client.post(url + 'confirm/', {'confirmation': preview.data['confirmation']}, format='json')
        self.assertEqual(confirmed.data['added'], 1)
        replay = self.client.post(url + 'confirm/', {'confirmation': preview.data['confirmation']}, format='json')
        self.assertEqual(replay.data['added'], 0)
        self.assertEqual(Student.objects.get().student_id, '000123')
        self.client.force_authenticate(self.other)
        room2 = Classroom.objects.create(owner=self.other, name='10B')
        self.assertEqual(self.client.post(f'/api/v1/classrooms/{room2.pk}/import/confirm/',
                                         {'confirmation': preview.data['confirmation']}, format='json').status_code, 400)

    def test_import_duplicate_missing_bad_file(self):
        url = f'/api/v1/classrooms/{self.room.pk}/import/preview/'
        first = self.client.post(url, {'file': SimpleUploadedFile('lop.csv', 'Họ tên,SBD\nAn,001\nBình,001\nThiếu,\n'.encode())})
        preview = self.client.post(url, {'source_token': first.data['source_token'], 'name_column': 0, 'id_column': 1}, format='json')
        self.assertEqual(len(preview.data['issues']), 2)
        self.assertIsNone(preview.data['confirmation'])
        self.assertEqual(self.client.post(url, {'file': SimpleUploadedFile('bad.xlsx', b'not an xlsx')}).status_code, 400)

    def test_reports_use_all_saved_scores_and_review_does_not_modify(self):
        self.student()
        self.student('000124', 'Chưa có bài')
        self.attach()
        old = self.submission(score=7)
        newer = self.submission(score=8)
        detail = {'part1_detail': {'1': {'student': 'A', 'correct': 'A', 'is_correct': True},
                                   '2': {'student': '?', 'correct': 'B', 'is_correct': False}},
                  'part2_detail': {'1': {k: {'student': 'Dung', 'correct': 'Dung', 'is_correct': True} for k in 'abcd'}}}
        for index in range(123):
            self.submission(str(index + 1000), score=6.5, detail_json=json.dumps(detail))
        data = report_data(self.exam)
        self.assertEqual(data['summary']['scored'], 124)
        self.assertEqual(data['summary']['missing'], 1)
        self.assertEqual(data['rows'][0]['score'], 8)
        self.assertIsNone(data['rows'][1]['score'])
        url = f'/api/v1/submissions/{old.pk}/review/'
        self.assertEqual(self.client.post(url, {'representative': True, 'reviewed': True}, format='json').status_code, 200)
        self.assertEqual(report_data(self.exam)['rows'][0]['score'], 7)
        old.refresh_from_db()
        newer.refresh_from_db()
        self.assertEqual((old.score, newer.score), (7, 8))
        workbook = load_workbook(io.BytesIO(xlsx_report(data)))
        self.assertEqual(workbook['Bảng điểm'].max_row, 126)
        self.assertEqual(workbook['Bảng điểm']['A2'].value, '000123')
        self.assertTrue(pdf_report(data).startswith(b'%PDF'))
        self.assertIn('Nguyễn An'.encode(), csv_report(data))
        question = next(q for q in data['questions'] if q['part'] == 1 and q['question'] == '2')
        self.assertEqual(question['unclear'], 123)
        self.assertEqual(question['wrong'], 0)
        self.assertIsNone(question['correct_rate'])
        part2 = [q for q in data['questions'] if q['part'] == 2]
        self.assertEqual(len(part2), 5)
        page = self.client.get('/api/v1/submissions/?limit=100&page=2')
        self.assertEqual(len(page.data['submissions']), 25)
        self.assertEqual(page.data['count'], 125)
        self.assertEqual(self.client.get('/api/v1/submissions/?limit=-1').status_code, 400)

    def test_manual_match_review_permissions_and_missing_details(self):
        self.student()
        self.attach()
        sub = self.submission('?', 5)
        roster = ExamRosterEntry.objects.get()
        response = self.client.post(f'/api/v1/submissions/{sub.pk}/review/', {'roster_id': roster.pk}, format='json')
        self.assertEqual(response.status_code, 200)
        data = report_data(self.exam)
        self.assertEqual(data['rows'][0]['submission_id'], sub.pk)
        self.assertEqual(data['summary']['missing_detail'], 1)
        sub.refresh_from_db()
        self.assertEqual(sub.student_id, '?')
        named = self.client.get('/api/v1/submissions/?q=Nguyễn')
        self.assertEqual(named.data['count'], 1)
        self.assertEqual(named.data['submissions'][0]['student_name'], 'Nguyễn An')
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.post(f'/api/v1/submissions/{sub.pk}/review/', {'reviewed': True}).status_code, 404)

    def test_unknown_identifiers_are_not_merged_into_one_student(self):
        self.submission('?', score=5)
        self.submission('?', score=6)
        self.submission('', score=7)
        data = report_data(self.exam)
        self.assertEqual(data['summary']['scored'], 3)
        self.assertTrue(all(row['duplicate_count'] == 1 for row in data['rows']))

    def test_manual_match_overrides_raw_sbd_for_class_and_name_filters(self):
        self.student()
        room2 = Classroom.objects.create(owner=self.owner, name='10B')
        Student.objects.create(classroom=room2, student_id='000124', name='Trần Bình')
        self.client.post(f'/api/v1/exams/{self.exam.pk}/roster/', {'classroom_ids': [self.room.pk, room2.pk]}, format='json')
        sub = self.submission()
        entry = ExamRosterEntry.objects.get(exam=self.exam, student_id='000124')
        self.client.post(f'/api/v1/submissions/{sub.pk}/review/', {'roster_id': entry.pk}, format='json')
        self.assertEqual(self.client.get(f'/api/v1/submissions/?classroom_id={self.room.pk}').data['count'], 0)
        response = self.client.get(f'/api/v1/submissions/?classroom_id={room2.pk}&q=Bình')
        self.assertEqual(response.data['count'], 1)
        self.assertEqual(response.data['submissions'][0]['student_name'], 'Trần Bình')
        self.assertEqual(self.client.get('/api/v1/submissions/?q=Nguyễn').data['count'], 0)

    def test_class_filter_and_roster_search_without_rewriting_raw_name(self):
        self.student()
        self.attach()
        sub = self.submission()
        self.submission('999999', score=8)
        response = self.client.get(f'/api/v1/submissions/?classroom_id={self.room.pk}&q=Nguyễn')
        self.assertEqual(response.data['count'], 1)
        self.assertEqual(response.data['submissions'][0]['class_name'], '10A')
        sub.refresh_from_db()
        self.assertEqual(sub.student_name, '')
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(f'/api/v1/submissions/?classroom_id={self.room.pk}').status_code, 404)

    def test_saved_details_variant_separation_and_formula_safety(self):
        for code, ok in [('101', True), ('202', False)]:
            variant = ExamVariant.objects.create(exam=self.exam, variant_code=code)
            self.submission(code, 6, variant=variant, student_name='=1+1', detail_json=json.dumps(
                {'part1_detail': {'1': {'student': 'A', 'correct': 'A' if ok else 'B', 'is_correct': ok}}}))
        data = report_data(self.exam)
        self.assertEqual(len(data['questions']), 2)
        self.assertEqual(sorted(q['correct_rate'] for q in data['questions']), [0, 100])
        workbook = load_workbook(io.BytesIO(xlsx_report(data)))
        self.assertEqual(workbook['Bảng điểm']['B2'].data_type, 's')
        self.assertTrue(workbook['Bảng điểm']['B2'].value.startswith("'="))

    def test_support_admin_image_privacy(self):
        with TemporaryDirectory() as media, override_settings(MEDIA_ROOT=media):
            from PIL import Image
            image = io.BytesIO()
            Image.new('RGB', (12, 12), 'white').save(image, 'PNG')
            response = self.client.post('/api/v1/support/', {'message': 'Có lỗi khi quét phiếu',
                'technical': json.dumps({'version': '1.0.1', 'token': 'must-not-save'}),
                'image': SimpleUploadedFile('photo.png', image.getvalue(), 'image/png')}, format='multipart')
            self.assertEqual(response.status_code, 201, response.data)
            ticket = SupportTicket.objects.get()
            self.assertNotIn('token', ticket.technical)
            self.assertEqual(self.client.get(f'/api/v1/support/{ticket.pk}/image/').status_code, 200)
            self.client.force_authenticate(self.other)
            self.assertEqual(self.client.get(f'/api/v1/support/{ticket.pk}/image/').status_code, 404)
            self.assertEqual(self.client.get('/api/v1/admin/support/').status_code, 403)
            self.other.is_superuser = True
            self.other.save()
            self.assertEqual(self.client.patch(f'/api/v1/admin/support/{ticket.pk}/', {'status': 'resolved'}, format='json').status_code, 200)

    def test_template_pdf_supported_and_unknown(self):
        from django.conf import settings
        if not (Path(settings.BASE_DIR) / 'cacmaubaithi').exists():
            self.skipTest('Template image fixtures not available')
        response = self.client.get('/api/v1/templates/40-08-06/pdf/')
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertEqual(self.client.get('/api/v1/templates/unknown/pdf/').status_code, 404)
